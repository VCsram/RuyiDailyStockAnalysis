# -*- coding: utf-8 -*-
"""Generate Word operation log for branding exercise 1."""
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[3]
OUT_DIR = ROOT / "docs" / "练习记录" / "练习1_品牌化改造"
SHOTS = OUT_DIR / "screenshots"
OUT_DOC = OUT_DIR / "练习1_品牌化改造_操作记录.docx"


def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    return p


def add_shot(doc, filename, caption):
    path = SHOTS / filename
    add_para(doc, caption, bold=True)
    if path.is_file():
        doc.add_picture(str(path), width=Cm(15.5))
        p = doc.add_paragraph(f"文件：{path.as_posix()}")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        add_para(doc, f"[缺失截图] {filename}")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    title = doc.add_heading("练习 1｜品牌化改造 — 操作记录", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_para(doc, f"仓库路径：D:\\quant\\RuyiDailyStockAnalysis（junction → D:\\tmp\\RuyiDailyStockAnalysis_v0.1.0）")
    add_para(doc, f"记录时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    add_para(doc, "品牌约定：中文名「如意金股」；英文名 RuyiDailyStockAnalysis；作者 creeper")

    add_heading(doc, "1. 准备", level=1)
    add_para(doc, "确认当前分支为 dev；git status 显示 No commits yet，全部为未跟踪文件。")
    add_para(doc, "策略：保留既有未跟踪内容，仅暂存并提交本练习相关文件。")

    add_heading(doc, "2. 改造范围（用户可见品牌）", level=1)
    for item in [
        "Web：index.html 标题/favicon、SidebarNav、LoginPage、uiText 中英标题与通知测试文案、brand.ts、ruyi-logo.png",
        "API：FastAPI title/description、RootResponse example、通知测试默认值、OpenAPI api_spec.json",
        "通知：EMAIL_SENDER_NAME 默认「如意金股分析助手」、system_config 默认测试标题/正文",
        "文档：README.md / README_EN.md / README_CHT.md / CHANGELOG [Unreleased] / AGENTS.md 对外品牌名",
        "桌面：loading.html 展示名（保留 package.json productName / Uninstall Daily Stock Analysis.exe 兼容路径）",
        "CLI：main.py 启动日志与 argparse description",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "3. 构建与启动命令", level=1)
    add_para(doc, "前端构建（须在真实路径，避免 junction 导致 Vite fileName 报错）：")
    add_para(doc, "cd D:\\tmp\\RuyiDailyStockAnalysis_v0.1.0\\apps\\dsa-web")
    add_para(doc, "npm run build  → 成功，产物写入 ../../static/")
    add_para(doc, "服务启动：")
    add_para(doc, ".\\.venv\\Scripts\\python.exe main.py --serve-only --host 127.0.0.1 --port 8000")
    add_para(doc, "说明：原 8000 进程复用后仍为旧 API 标题，已重启以加载新 OpenAPI 品牌。")

    add_heading(doc, "4. 健康检查", level=1)
    add_para(doc, "GET http://127.0.0.1:8000/api/health → HTTP 200，body 含 status=ok")
    add_para(doc, '示例：{"status":"ok","timestamp":"2026-07-17T11:02:37.946031"}')
    add_para(doc, "GET http://127.0.0.1:8000/ → HTTP 200，index.html 含 /ruyi-logo.png 与标题「如意金股」")
    add_para(doc, "GET http://127.0.0.1:8000/openapi.json → title = RuyiDailyStockAnalysis API")

    add_heading(doc, "5. 真实页面截图", level=1)
    add_shot(doc, "01_home_sidebar.png", "图 1：首页打开侧边栏，可见 Logo「如意」与文字品牌「如意金股」")
    add_shot(doc, "01_home_sidebar_pw.png", "图 2：Playwright 补充截图（首页侧栏）")
    add_shot(doc, "02_login_pw.png", "图 3：/login 路由（当前管理员认证关闭时会回首页；代码层 LoginPage 已改为如意金股 + RuyiDailyStockAnalysis + Author creeper）")
    add_shot(doc, "03_api_docs.png", "图 4：/docs Swagger 标题与描述已为 RuyiDailyStockAnalysis / 如意金股")

    add_heading(doc, "6. 旧品牌残留检查", level=1)
    add_para(doc, "对 apps/dsa-web/src、API、README、通知默认值等展示层扫描：")
    doc.add_paragraph("creeper.png / <title>creeper> / DSA 通知测试 / daily_stock_analysis股票分析助手 / 如意每日金股分析系统 → 展示层已清理", style="List Bullet")
    doc.add_paragraph("保留：GitHub Actions 工作流显示名「Daily Stock Analysis」、Docker 镜像 zhulinsen/daily_stock_analysis、桌面 Uninstall Daily Stock Analysis.exe 兼容探测、上游仓库 URL", style="List Bullet")
    doc.add_paragraph("作者 creeper 仅作为 Author 展示保留（Login 页脚 / README）", style="List Bullet")

    add_heading(doc, "7. 提交", level=1)
    add_para(doc, "消息：feat: 将项目品牌统一为如意金股")
    add_para(doc, "仅暂存本练习相关源文件与本操作记录；不推送 remote。")
    add_para(doc, "（提交哈希见文末，由提交后脚本回填或手工填写）")

    add_heading(doc, "8. 边界遵守说明", level=1)
    for item in [
        "未改接口路径与字段契约（仅改展示文案/默认通知标题）",
        "未改数据库、依赖、密钥",
        "未手改 static/ 构建产物（通过 npm run build 生成）",
        "未改桌面 appId / productName，避免破坏安装目录与卸载兼容路径",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(OUT_DOC)
    print("wrote", OUT_DOC)


if __name__ == "__main__":
    main()
