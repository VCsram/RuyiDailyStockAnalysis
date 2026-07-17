# -*- coding: utf-8 -*-
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[3]
OUT_DIR = ROOT / "docs" / "练习记录" / "练习2_Logo与Icon"
SHOTS = OUT_DIR / "screenshots"
OUT_DOC = OUT_DIR / "练习2_Logo与Icon_操作记录.docx"


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_shot(doc, name, caption):
    add_para(doc, caption, bold=True)
    path = SHOTS / name
    if path.is_file():
        doc.add_picture(str(path), width=Cm(15.2))
        p = doc.add_paragraph(str(path.as_posix()))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        add_para(doc, f"[缺失] {name}")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    # copy any MCP screenshots dropped in temp
    import shutil

    tmp = Path.home() / "AppData/Local/Temp/cursor/screenshots/docs/练习记录/练习2_Logo与Icon/screenshots"
    if tmp.is_dir():
        for p in tmp.glob("*.png"):
            shutil.copy2(p, SHOTS / p.name)

    doc = Document()
    t = doc.add_heading("练习 2｜设计并实现项目 Logo 与 Icon — 操作记录", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, f"仓库：D:\\quant\\RuyiDailyStockAnalysis（真实路径 D:\\tmp\\RuyiDailyStockAnalysis_v0.1.0）")
    add_para(doc, f"记录时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    add_para(doc, "品牌：如意金股 / RuyiDailyStockAnalysis；作者 creeper（未改名）")

    doc.add_heading("1. 设计取舍", level=1)
    for item in [
        "方向 A：如意结缠绕 K 线 — 文化强，16px 易糊。",
        "方向 B：几何「如」字 + AI 节点 — 跨语言辨识弱。",
        "方向 C（选定）：玉如意弧 + 上升折线 + AI 节点 — 同时表达如意/走势/量化，粗线+节点在 16px 可辨。",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("2. 资产路径", level=1)
    for item in [
        "源文件：docs/assets/dsa_vi/{icon,logo-light,logo-dark,favicon}.svg + favicon.ico + png/icon-*.png + README.md",
        "生成脚本：docs/assets/dsa_vi/build_brand_assets.py",
        "Web 副本：apps/dsa-web/public/brand/ 与 public/favicon.svg|ico",
        "登录视觉预览（不改 .env）：/brand/login-preview.html",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3. 代码位置", level=1)
    for item in [
        "BrandMark.tsx → apps/dsa-web/src/components/brand/BrandMark.tsx",
        "SidebarNav.tsx / LoginPage.tsx 接入 BrandMark",
        "index.html 更新 favicon.ico + favicon.svg",
        "brand.ts 指向 /brand/*.svg",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("4. 验证", level=1)
    add_para(doc, "vitest：BrandMark.test.tsx + SidebarNav.test.tsx → 2 files / 13 tests passed")
    add_para(doc, "npm run build（真实路径 apps/dsa-web）→ success，产物含 static/brand 与 favicon")
    add_para(doc, "资产自检：SVG ElementTree 可解析；PNG 尺寸 16..512 正确；ICO entries=6")
    add_para(doc, 'GET /api/health → 200 {"status":"ok",...}')
    add_para(doc, "GET / → 200；GET /brand/icon.svg → 200；GET /favicon.ico → 200")

    doc.add_heading("5. 截图", level=1)
    add_shot(doc, "01_home_sidebar.png", "图 1：主界面打开侧栏（深色）")
    add_shot(doc, "01b_sidebar_brandmark.png", "图 2：侧栏 BrandMark（图标 + 如意金股）")
    add_shot(doc, "02_home_sidebar_light.png", "图 3：切换浅色主题后侧栏")
    add_shot(doc, "03_login_preview.png", "图 4：登录页 Logo 视觉预览（认证关闭，不改 .env）")
    add_shot(doc, "03b_login_preview_mcp.png", "图 5：登录预览补充截图")
    add_shot(doc, "04_favicon_svg.png", "图 6：favicon.svg")

    doc.add_heading("6. Git", level=1)
    add_para(doc, "本次按要求：不提交、不推送。保留用户既有未跟踪内容。")

    doc.save(OUT_DOC)
    print("wrote", OUT_DOC)


if __name__ == "__main__":
    main()
